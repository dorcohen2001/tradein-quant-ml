"""Persistent interval gates, retention sweeps, chart specs, and store I/O."""
from __future__ import annotations

import datetime as dt
import json
import os
import time

import pytest

from core.charting.spec import ChartSpec
from core.data.cache_policy import bucket_seconds_for
from core.maintenance import prune_by_age, sweep_runtime
from core.paper_trading.store import load_trades, pnl_pct, save_trades
from core.scheduling import IntervalGate, PersistentIntervalGate


# --------------------------------------------------------------------------
# Interval gates
# --------------------------------------------------------------------------
def test_in_memory_gate_fires_immediately_then_throttles():
    gate = IntervalGate(15)
    start = dt.datetime(2026, 8, 15, 12, 0)
    assert gate.due(start) is True
    assert gate.due(start + dt.timedelta(minutes=14)) is False
    assert gate.due(start + dt.timedelta(minutes=15)) is True


def test_persistent_gate_survives_process_restarts(tmp_path):
    """The reason this class exists: under `--once` scheduling the runner is a
    fresh process every tick, so an in-memory marker would reset each time and
    the gate would degrade to 'fire every cycle'. A NEW gate object each call
    stands in for a new process."""
    path = str(tmp_path / "gates.json")
    start = dt.datetime(2026, 8, 15, 12, 0)

    def tick(minutes):
        return PersistentIntervalGate(15, state_path=path, name="funnel").due(
            start + dt.timedelta(minutes=minutes))

    fired = [m for m in range(0, 18, 2) if tick(m)]
    assert fired == [0, 16]  # exactly twice across a 2-minute cadence


def test_persistent_gates_are_namespaced_within_one_file(tmp_path):
    path = str(tmp_path / "gates.json")
    now = dt.datetime(2026, 8, 15, 12, 0)
    assert PersistentIntervalGate(15, path, "a").due(now) is True
    assert PersistentIntervalGate(15, path, "b").due(now) is True   # independent
    assert PersistentIntervalGate(15, path, "a").due(now) is False
    assert set(json.loads(open(path).read())) == {"a", "b"}


def test_corrupt_gate_state_opens_rather_than_jamming_shut(tmp_path):
    """A throttle stuck open is a nuisance; stuck shut silently disables the
    work it guards, which is far harder to notice."""
    path = tmp_path / "gates.json"
    path.write_text("{ not json", encoding="utf-8")
    assert PersistentIntervalGate(15, str(path), "f").due() is True


def test_unwritable_gate_state_still_fires(tmp_path, monkeypatch):
    """Simulated by making the write raise, rather than by naming a bogus path:
    on Windows a path like "/nonexistent-drive/x" is simply resolved against
    the current drive and cheerfully CREATED, so that version of this test both
    passed for the wrong reason and littered the filesystem."""
    gate = PersistentIntervalGate(15, state_path=str(tmp_path / "g.json"), name="f")
    monkeypatch.setattr(
        gate, "_write_state",
        lambda state: (_ for _ in ()).throw(OSError("read-only filesystem")),
    )
    assert gate.due() is True


def test_gate_state_write_is_atomic(tmp_path):
    path = str(tmp_path / "gates.json")
    PersistentIntervalGate(15, path, "f").due()
    assert [p.name for p in tmp_path.iterdir()] == ["gates.json"]


# --------------------------------------------------------------------------
# Retention
# --------------------------------------------------------------------------
def _aged(path, days):
    path.write_bytes(b"x" * 100)
    old = time.time() - days * 86400
    os.utime(path, (old, old))


def test_prune_removes_only_files_past_the_age(tmp_path):
    _aged(tmp_path / "old.pkl", 30)
    _aged(tmp_path / "new.pkl", 1)
    removed, reclaimed = prune_by_age(str(tmp_path), max_age_days=7, suffix=".pkl")
    assert removed == 1 and reclaimed == 100
    assert {p.name for p in tmp_path.iterdir()} == {"new.pkl"}


def test_prune_respects_the_suffix_filter(tmp_path):
    _aged(tmp_path / "old.pkl", 30)
    _aged(tmp_path / "keep.json", 30)
    prune_by_age(str(tmp_path), max_age_days=7, suffix=".pkl")
    assert {p.name for p in tmp_path.iterdir()} == {"keep.json"}


def test_zero_retention_disables_the_sweep(tmp_path):
    """A misread config must never mean 'delete everything'."""
    _aged(tmp_path / "ancient.pkl", 9999)
    assert prune_by_age(str(tmp_path), max_age_days=0, suffix=".pkl") == (0, 0)
    assert (tmp_path / "ancient.pkl").exists()


def test_dry_run_reports_without_deleting(tmp_path):
    _aged(tmp_path / "old.pkl", 30)
    removed, _ = prune_by_age(str(tmp_path), max_age_days=7, suffix=".pkl", dry_run=True)
    assert removed == 1
    assert (tmp_path / "old.pkl").exists()


def test_prune_tolerates_a_missing_directory():
    assert prune_by_age("/no/such/dir", max_age_days=7) == (0, 0)


def test_sweep_targets_only_cache_and_charts(tmp_path):
    (tmp_path / "cache" / "yfinance").mkdir(parents=True)
    (tmp_path / "charts").mkdir()
    (tmp_path / "stores").mkdir()
    _aged(tmp_path / "cache" / "yfinance" / "a.pkl", 30)
    _aged(tmp_path / "charts" / "b.png", 30)
    _aged(tmp_path / "stores" / "trades_livne.json", 9999)  # must NEVER be touched

    results = sweep_runtime(str(tmp_path), cache_days=7, chart_days=14)
    assert results["yfinance_cache"][0] == 1
    assert results["charts"][0] == 1
    assert (tmp_path / "stores" / "trades_livne.json").exists()


# --------------------------------------------------------------------------
# Store
# --------------------------------------------------------------------------
def test_store_round_trip(tmp_path):
    path = str(tmp_path / "trades.json")
    trades = [{"id": "a", "current_status": "ACTIVE"}]
    save_trades(path, trades)
    assert load_trades(path) == trades


def test_missing_store_is_an_empty_sample(tmp_path):
    assert load_trades(str(tmp_path / "none.json")) == []


def test_corrupt_store_raises_instead_of_returning_empty(tmp_path):
    """Returning [] would let the next save overwrite the file and erase all
    history. This is the single most important line in the store module."""
    path = tmp_path / "trades.json"
    path.write_text("{ truncated", encoding="utf-8")
    with pytest.raises(RuntimeError, match="corrupt"):
        load_trades(str(path))


def test_migration_hook_is_applied_per_record(tmp_path):
    path = str(tmp_path / "trades.json")
    save_trades(path, [{"id": "a"}, {"id": "b"}])
    out = load_trades(path, migrate=lambda t: {**t, "migrated": True})
    assert all(t["migrated"] for t in out)


def test_save_leaves_no_temp_files(tmp_path):
    path = str(tmp_path / "trades.json")
    save_trades(path, [{"id": "a"}])
    save_trades(path, [{"id": "b"}])
    assert [p.name for p in tmp_path.iterdir()] == ["trades.json"]


# --------------------------------------------------------------------------
# Chart spec
# --------------------------------------------------------------------------
def test_daily_spec_reproduces_the_original_hardcoded_chart():
    spec = ChartSpec.daily()
    assert (spec.interval, spec.tf_label, spec.stock_period, spec.crypto_limit) == \
           ("1d", "1D", "2y", 600)
    assert spec.emas == (50, 200) and spec.show_rsi and spec.show_macd
    assert spec.warnings() == []


def test_intraday_spec_stays_inside_the_yfinance_cap():
    """Asking yfinance for 2y of 15m returns an EMPTY frame, which would
    degrade silently to text-only alerts."""
    spec = ChartSpec.intraday_15m()
    assert spec.interval == "15m"
    assert spec.warnings() == []


def test_spec_warns_when_history_exceeds_the_intraday_cap():
    spec = ChartSpec(interval="15m", tf_label="15M", stock_period="2y")
    issues = spec.warnings()
    assert issues and "cap" in issues[0]


def test_spec_warns_on_insufficient_ema_warmup():
    spec = ChartSpec(interval="1d", crypto_limit=50, emas=(200,))
    assert any("warm-up" in w for w in spec.warnings())


def test_spec_is_frozen():
    with pytest.raises(Exception):
        ChartSpec.daily().interval = "5m"


# --------------------------------------------------------------------------
# Cache policy
# --------------------------------------------------------------------------
@pytest.mark.parametrize("interval,expected", [
    ("5m", 60), ("15m", 60), ("30m", 60),
    ("1h", 300), ("4h", 300),
    ("1d", 1800), ("1w", 1800),
])
def test_cache_bucket_never_outlives_the_bar(interval, expected):
    """Biton previously cached a 1h frame for 30 minutes -- half the bar
    period -- so every second cycle re-used the previous cycle's bars."""
    assert bucket_seconds_for(interval) == expected


# --------------------------------------------------------------------------
# Journal resilience -- a locked spreadsheet must never cost an alert
# --------------------------------------------------------------------------
def test_journal_write_round_trips(tmp_path):
    from core.paper_trading.journal import upsert_row
    path = str(tmp_path / "j.xlsx")
    cols = ["ID", "Ticker", "Status"]
    assert upsert_row(path, cols, ["a1", "BTCUSDT", "ACTIVE"], "a1") is True
    assert upsert_row(path, cols, ["a1", "BTCUSDT", "WIN_TP2"], "a1") is True  # upsert, not append

    from openpyxl import load_workbook
    rows = list(load_workbook(path).active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][2] == "WIN_TP2"


# Scoped and deliberate: openpyxl builds its OWN ZipFile to write with, and when
# that write fails against a locked file the object is finalized later during GC,
# where CPython reports the close error as unraisable. It is not our exception --
# `upsert_row` catches and reports the real one, which is what this test asserts.
@pytest.mark.filterwarnings("ignore::pytest.PytestUnraisableExceptionWarning")
def test_locked_journal_returns_false_and_never_raises(tmp_path):
    """THE regression guard. Before this, openpyxl's PermissionError propagated
    out of update_open_positions AFTER the trade store was written but BEFORE
    the caller sent its Telegram alerts -- so the transition persisted, the
    alert was never sent, and the trade was no longer in `changed` next cycle.
    The alert was lost permanently because a spreadsheet was open in Excel."""
    import msvcrt

    from core.paper_trading.journal import upsert_row
    path = str(tmp_path / "j.xlsx")
    cols = ["ID", "Ticker", "Status"]
    upsert_row(path, cols, ["a1", "BTCUSDT", "ACTIVE"], "a1")

    handle = open(path, "r+b")
    try:
        msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
    except OSError:
        pytest.skip("could not acquire an exclusive lock on this platform")
    try:
        assert upsert_row(path, cols, ["a2", "ETHUSDT", "LOSS_SL"], "a2") is False
    finally:
        try:
            msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        except OSError:
            pass
        handle.close()


def test_journal_failure_on_a_bad_path_is_non_fatal(tmp_path):
    from core.paper_trading.journal import upsert_row
    assert upsert_row(str(tmp_path / "\0bad" / "j.xlsx"), ["ID"], ["x"], "x") is False


# --------------------------------------------------------------------------
# CandleCache concurrency
# --------------------------------------------------------------------------
def test_candle_cache_temp_file_is_pid_scoped(tmp_path, monkeypatch):
    """Every tenant runner is its own process sharing this directory. A fixed
    `<file>.tmp` meant two tenants fetching the same symbol raced on one path:
    a PermissionError on replace, or one truncating the other's partial write."""
    import pandas as pd

    from core.backtesting.cache import CandleCache

    seen = []
    real_to_csv = pd.DataFrame.to_csv

    def spy(self, path_or_buf=None, *a, **k):
        if isinstance(path_or_buf, (str, os.PathLike)):
            seen.append(str(path_or_buf))
        return real_to_csv(self, path_or_buf, *a, **k)

    monkeypatch.setattr(pd.DataFrame, "to_csv", spy)

    cache = CandleCache(str(tmp_path), fmt="csv")
    frame = pd.DataFrame(
        {"Open": [1.0], "High": [2.0], "Low": [0.5], "Close": [1.5], "Volume": [10]},
        index=pd.DatetimeIndex([dt.datetime(2026, 8, 15, 12, 0)]),
    )
    out = cache.save("binance", "BTCUSDT", "15m", frame)

    assert seen, "expected a temp write"
    assert str(os.getpid()) in seen[0], f"temp name is not PID-scoped: {seen[0]}"
    assert out.exists()
    assert not list(tmp_path.rglob("*.tmp")), "a scratch file was left behind"


def test_candle_cache_save_is_readable_back(tmp_path):
    import pandas as pd

    from core.backtesting.cache import CandleCache
    cache = CandleCache(str(tmp_path), fmt="csv")
    frame = pd.DataFrame(
        {"Open": [1.0], "High": [2.0], "Low": [0.5], "Close": [1.5], "Volume": [10]},
        index=pd.DatetimeIndex([dt.datetime(2026, 8, 15, 12, 0)]),
    )
    cache.save("binance", "BTCUSDT", "15m", frame)
    assert len(cache.load("binance", "BTCUSDT", "15m")) == 1


# --------------------------------------------------------------------------
# Candle cache retention
# --------------------------------------------------------------------------
def test_sweep_prunes_the_candle_cache_by_its_own_age(tmp_path):
    """`data/` grew unbounded before -- it is written by every tenant on every
    cycle and no sweep covered it. It gets a LONGER retention than the yfinance
    cache because backtest histories are slow and rate-limited to rebuild."""
    runtime = tmp_path / "runtime"
    (runtime / "cache" / "yfinance").mkdir(parents=True)
    (runtime / "charts").mkdir()
    candles = tmp_path / "data" / "binance" / "BTCUSDT"
    candles.mkdir(parents=True)

    _aged(candles / "15m.parquet", 45)   # older than the 30-day candle window
    _aged(candles / "1h.parquet", 10)    # inside it
    _aged(runtime / "cache" / "yfinance" / "x.pkl", 10)  # older than the 7-day window

    results = sweep_runtime(str(runtime), cache_days=7, chart_days=14,
                            candle_dir=str(tmp_path / "data"), candle_days=30)

    assert results["candle_cache"][0] == 1
    assert not (candles / "15m.parquet").exists()
    assert (candles / "1h.parquet").exists()
    assert results["yfinance_cache"][0] == 1


def test_sweep_without_a_candle_dir_is_unchanged(tmp_path):
    """Back-compat: existing callers that pass no candle dir must behave
    exactly as before, with no candle_cache key in the result."""
    runtime = tmp_path / "runtime"
    (runtime / "charts").mkdir(parents=True)
    results = sweep_runtime(str(runtime), cache_days=7, chart_days=14)
    assert "candle_cache" not in results


# --------------------------------------------------------------------------
# Shared market-data layer
# --------------------------------------------------------------------------
def test_interval_minutes_parsing():
    from core.data.market import interval_minutes
    assert interval_minutes("15m") == 15
    assert interval_minutes("1h") == 60
    assert interval_minutes("1d") == 1440
    with pytest.raises(ValueError):
        interval_minutes("nonsense")


def test_window_end_is_quantized_so_tenants_request_the_same_range():
    """THE deduplication mechanism. Two tenant processes calling seconds apart
    must floor to the SAME instant, or the second sees a 'missing' trailing
    segment and downloads it."""
    from core.data.market import quantized_end
    a = quantized_end("15m", dt.datetime(2026, 8, 15, 12, 0, 17))
    b = quantized_end("15m", dt.datetime(2026, 8, 15, 12, 0, 43))
    assert a == b
    # ...and a bucket later it does advance, so data never goes stale.
    later = quantized_end("15m", dt.datetime(2026, 8, 15, 12, 5, 0))
    assert later > a


def test_second_tenant_reuses_the_cache_with_zero_downloads():
    """Model A fetches BTCUSDT 15m; Model B asking for the same symbol and
    interval inside the cache window must cost no network call."""
    import pandas as pd

    from core.data.market import get_candles

    class StubFeed:
        """Stands in for IntradayDataFeed: serves from a dict, counting only
        genuine cache misses as downloads."""
        def __init__(self):
            self.download_count = 0
            self._cache = {}

        def get_candles(self, symbol, asset_class, interval, start, end):
            key = (symbol, asset_class, interval, start, end)
            if key not in self._cache:
                self.download_count += 1
                idx = pd.date_range(end - dt.timedelta(minutes=15 * 5), periods=5, freq="15min")
                self._cache[key] = pd.DataFrame(
                    {"Open": 1.0, "High": 2.0, "Low": 0.5, "Close": 1.5, "Volume": 10}, index=idx)
            return self._cache[key]

    feed = StubFeed()
    now = dt.datetime(2026, 8, 15, 12, 0, 5)

    get_candles("BTCUSDT", "crypto", "15m", bars=5, now=now, feed=feed)          # tenant A
    assert feed.download_count == 1
    get_candles("BTCUSDT", "crypto", "15m", bars=5,
                now=now + dt.timedelta(seconds=30), feed=feed)                    # tenant B
    assert feed.download_count == 1, "tenant B should have hit the shared cache"

    # A different symbol is genuinely a new fetch.
    get_candles("ETHUSDT", "crypto", "15m", bars=5, now=now, feed=feed)
    assert feed.download_count == 2


def test_shared_layer_returns_the_canonical_frame_shape():
    """Every consumer in this repo expects OHLCV columns on a naive-UTC index,
    oldest first -- the shape core.backtesting.cache.normalize_frame produces."""
    import pandas as pd

    from core.data.market import get_candles

    class StubFeed:
        download_count = 0

        def get_candles(self, symbol, asset_class, interval, start, end):
            idx = pd.date_range(end - dt.timedelta(minutes=15 * 10), periods=10, freq="15min")
            return pd.DataFrame(
                {"Open": 1.0, "High": 2.0, "Low": 0.5, "Close": 1.5, "Volume": 10}, index=idx)

    frame = get_candles("BTCUSDT", "crypto", "15m", bars=4, feed=StubFeed())
    assert list(frame.columns) == ["Open", "High", "Low", "Close", "Volume"]
    assert frame.index.tz is None
    assert len(frame) == 4                      # tail-sliced to `bars`
    assert frame.index.is_monotonic_increasing  # oldest first


def test_empty_result_is_returned_not_raised():
    """One unlisted pair must skip, not abort a whole watchlist scan."""
    import pandas as pd

    from core.data.market import get_candles

    class EmptyFeed:
        download_count = 0

        def get_candles(self, *a, **k):
            return pd.DataFrame()

    assert len(get_candles("NOPEUSDT", "crypto", "15m", feed=EmptyFeed())) == 0


# --------------------------------------------------------------------------
# Sweep safety: ephemeral only, never persistent state
# --------------------------------------------------------------------------
def test_sweep_never_touches_stores_journals_or_backtest_logs(tmp_path):
    """The non-negotiable rule: PnL history and audit trails are permanent.
    Every one of these is ancient and must survive a full sweep."""
    runtime = tmp_path / "runtime"
    for sub in ("stores", "journals", "backtest", "logs", "charts", "funnel_logs"):
        (runtime / sub).mkdir(parents=True)

    permanent = [
        runtime / "stores" / "trades_livne.json",
        runtime / "stores" / "trades_client_alpha.json",
        runtime / "stores" / "report_state_livne.json",
        runtime / "journals" / "trade_journal_livne.xlsx",
        runtime / "backtest" / "trades_log_livne.csv",
    ]
    for path in permanent:
        _aged(path, 9999)

    ephemeral = [
        runtime / "charts" / "livne_BTCUSDT_15M_2020-01-01.png",
        runtime / "logs" / "livne.log.1",
        runtime / "funnel_logs" / "live_funnel_livne_20200101_000000.xlsx",
    ]
    for path in ephemeral:
        _aged(path, 9999)

    live_log = runtime / "logs" / "livne.log"
    _aged(live_log, 9999)  # ancient, but it is the LIVE log

    sweep_runtime(str(runtime), cache_days=7, chart_days=14, log_days=7, funnel_days=14)

    for path in permanent:
        assert path.exists(), f"PERMANENT STATE DELETED: {path.name}"
    for path in ephemeral:
        assert not path.exists(), f"ephemeral file survived: {path.name}"
    assert live_log.exists(), "the live log must be size-rotated, never aged out"
