"""Excel trade journal (`paper_trades_journal.xlsx`).

Mirrors the JSON paper-trading store (`paper_trading/tracker.py`) into a
human-readable spreadsheet so the trade history can be reviewed offline.
Rows are matched and updated by trade ID every time a trade is logged or its
status changes, so the sheet stays in sync with `trades.json` rather than
being appended to indefinitely.
"""
from __future__ import annotations

import os

from openpyxl import Workbook, load_workbook

from config import PAPER_TRADING_JOURNAL_XLSX

COLUMNS = [
    "ID", "Signal Date", "Ticker", "Asset Class", "Direction",
    "Entry Zone", "Stop Loss", "Initial Stop", "Target 1", "Target 2", "Status",
    "TP1 Hit", "TP1 Exit Price", "Breakeven Price", "Remaining Size %",
    "Emergency Exit Reason", "Telegram Message ID", "Exit Date", "Exit Price", "PnL %",
]


def _entry_zone_str(trade: dict) -> str:
    low, high = trade["entry_zone"]
    return f"{low:g} - {high:g}"


def _row_values(trade: dict) -> list:
    return [
        trade["id"],
        trade["signal_date"],
        trade["ticker"],
        trade["asset_class"],
        trade["direction"],
        _entry_zone_str(trade),
        trade["stop_loss"],
        trade.get("initial_stop_loss"),
        trade["target1"],
        trade["target2"],
        trade["current_status"],
        trade.get("tp1_hit", False),
        trade.get("tp1_exit_price"),
        trade.get("breakeven_price"),
        trade.get("remaining_size_pct"),
        trade.get("emergency_exit_reason"),
        trade.get("telegram_message_id"),
        trade.get("exit_date"),
        trade.get("exit_price"),
        trade.get("pnl_pct"),
    ]


def _open_or_create_workbook():
    if os.path.exists(PAPER_TRADING_JOURNAL_XLSX):
        wb = load_workbook(PAPER_TRADING_JOURNAL_XLSX)
        return wb, wb.active

    os.makedirs(os.path.dirname(PAPER_TRADING_JOURNAL_XLSX), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Trades"
    ws.append(COLUMNS)
    return wb, ws


def upsert_trade_row(trade: dict) -> None:
    """Insert or update the spreadsheet row for `trade`, matched by ID."""
    wb, ws = _open_or_create_workbook()

    trade_id = trade["id"]
    target_row = None
    for (id_cell,) in ws.iter_rows(min_row=2, max_col=1):
        if id_cell.value == trade_id:
            target_row = id_cell.row
            break

    values = _row_values(trade)
    if target_row is None:
        ws.append(values)
    else:
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=target_row, column=col_idx, value=value)

    wb.save(PAPER_TRADING_JOURNAL_XLSX)
